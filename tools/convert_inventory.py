# -*- coding: utf-8 -*-
"""Convert 財產清冊.xlsx sheet 0910 into data/inventory.json."""

from __future__ import annotations

import json
import sys
from collections import Counter
from datetime import date
from pathlib import Path

try:
    import openpyxl
except ImportError:
    sys.stderr.write("請先安裝 openpyxl：py -3 -m pip install openpyxl\n")
    sys.exit(1)

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "財產清冊.xlsx"
OUTPUT = ROOT / "data" / "inventory.json"


def cell_str(value) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        if value.is_integer():
            return str(int(value))
        return format(value, "f").rstrip("0").rstrip(".")
    return str(value).strip()


def roc_to_iso(value) -> str | None:
    raw = cell_str(value)
    if not raw:
        return None
    digits = "".join(ch for ch in raw if ch.isdigit())
    if len(digits) == 6:
        year, month, day = int(digits[:2]), int(digits[2:4]), int(digits[4:6])
    elif len(digits) == 7:
        year, month, day = int(digits[:3]), int(digits[3:5]), int(digits[5:7])
    else:
        return None
    try:
        converted = date(year + 1911, month, day)
    except ValueError:
        return None
    return converted.isoformat()


def to_number(value):
    if value is None or value == "":
        return None
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        if isinstance(value, float) and value.is_integer():
            return int(value)
        return int(value) if isinstance(value, int) else value
    text = cell_str(value).replace(",", "")
    if not text:
        return None
    try:
        number = float(text)
    except ValueError:
        return None
    return int(number) if number.is_integer() else number


def main() -> int:
    if not SOURCE.exists():
        sys.stderr.write(f"找不到清冊檔案：{SOURCE}\n")
        return 1

    workbook = openpyxl.load_workbook(SOURCE, data_only=True)
    if "0910" not in workbook.sheetnames:
        sys.stderr.write(f"找不到工作表 0910，現有：{workbook.sheetnames}\n")
        return 1

    sheet = workbook["0910"]
    items = []
    ids = []
    date_failures = []

    for row_number, row in enumerate(
        sheet.iter_rows(min_row=2, max_row=sheet.max_row, max_col=21, values_only=True),
        start=2,
    ):
        if all(cell is None or cell_str(cell) == "" for cell in row):
            continue

        property_id = cell_str(row[0])
        if not property_id:
            continue

        purchase_date = roc_to_iso(row[1])
        if row[1] not in (None, "") and purchase_date is None:
            date_failures.append({"row": row_number, "value": cell_str(row[1])})

        item = {
            "propertyId": property_id,
            "name": cell_str(row[2]),
            "location": cell_str(row[8]),
            "department": cell_str(row[7]),
            "custodian": cell_str(row[9]),
            "specification": cell_str(row[4]),
            "unit": cell_str(row[5]),
            "price": to_number(row[6]),
            "purchaseDate": purchase_date,
            "serviceLife": to_number(row[3]),
            "supplier": cell_str(row[11]),
            "status": cell_str(row[12]),
            "note": cell_str(row[10]),
            "brand": cell_str(row[19]),
            "model": cell_str(row[20]),
        }
        items.append(item)
        ids.append(property_id)

    duplicate_ids = sorted([pid for pid, count in Counter(ids).items() if count > 1])
    payload = {
        "source": "財產清冊.xlsx",
        "sheet": "0910",
        "count": len(items),
        "items": items,
    }

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    OUTPUT.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")

    report = {
        "count": len(items),
        "uniqueIds": len(set(ids)),
        "duplicateIds": duplicate_ids,
        "dateFailures": date_failures,
        "emptyNames": sum(1 for item in items if not item["name"]),
        "output": str(OUTPUT),
    }
    print(json.dumps(report, ensure_ascii=False, indent=2))
    return 0 if len(items) and not duplicate_ids and not date_failures else 1


if __name__ == "__main__":
    raise SystemExit(main())
